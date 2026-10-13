"""Compara dos versiones de un Word y genera el Excel de registro de cambios.

Hojas: "Registro de cambios" (ID, sección, tipo, antes, después, resaltado)
y "Resumen" (conteos por tipo y por sección).

Uso:
    python herramientas/comparar.py antes.docx despues.docx registro.xlsx
"""
import collections
import difflib
import sys

import docx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from docx_lib import W, parrafos


def bloques(ruta):
    out = []
    for p in parrafos(docx.Document(ruta)):
        t = p['texto'].strip()
        if not t:
            continue
        resaltado = sorted({h.get(W('val')) for h in p['el'].iter(W('highlight'))})
        out.append({'t': t, 'sec': p['seccion'], 'tabla': p['en_tabla'], 'hl': ','.join(resaltado)})
    return out


def main(antes, despues, salida):
    A, B = bloques(antes), bloques(despues)
    sm = difflib.SequenceMatcher(None, [x['t'] for x in A], [x['t'] for x in B], autojunk=False)
    filas = []
    for op, a1, a2, b1, b2 in sm.get_opcodes():
        if op == 'equal':
            continue
        if op == 'replace':
            for k in range(max(a2 - a1, b2 - b1)):
                xa = A[a1 + k] if a1 + k < a2 else None
                xb = B[b1 + k] if b1 + k < b2 else None
                tipo = 'Modificado' if xa and xb else ('Eliminado' if xa else 'Agregado')
                ref = xb or xa
                filas.append([ref['sec'], 'Tabla' if ref['tabla'] else 'Texto', tipo,
                              xa['t'] if xa else '', xb['t'] if xb else '', (xb or {}).get('hl', '')])
        elif op == 'delete':
            for xa in A[a1:a2]:
                filas.append([xa['sec'], 'Tabla' if xa['tabla'] else 'Texto', 'Eliminado', xa['t'], '', ''])
        elif op == 'insert':
            for xb in B[b1:b2]:
                filas.append([xb['sec'], 'Tabla' if xb['tabla'] else 'Texto', 'Agregado', '', xb['t'], xb['hl']])

    wb = Workbook()
    ws = wb.active
    ws.title = 'Registro de cambios'
    ws.append(['ID', 'Sección', 'Ubicación', 'Tipo de cambio', 'Antes', 'Después', 'Resaltado'])
    for n, f in enumerate(filas, 1):
        ws.append([f'C{n:02d}'] + [_seguro(x) for x in f])
    for col, ancho in zip('ABCDEFG', [7, 35, 10, 14, 70, 70, 12]):
        ws.column_dimensions[col].width = ancho
    r = wb.create_sheet('Resumen')
    r.append(['Concepto', 'Cantidad'])
    for k, v in collections.Counter(f[2] for f in filas).most_common():
        r.append([k, v])
    r.append([])
    r.append(['Sección', 'Cambios'])
    for k, v in collections.Counter(f[0] for f in filas).most_common():
        r.append([k or '(sin sección)', v])
    for hoja in wb:
        for c in hoja[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F4E78')
        hoja.freeze_panes = 'A2'
        for fila in hoja.iter_rows(min_row=2):
            for c in fila:
                c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(salida)
    print(f'cambios {len(filas)} -> {salida}')


def _seguro(v):
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@'):
        return "'" + v
    return v


if __name__ == '__main__':
    main(*sys.argv[1:4])
