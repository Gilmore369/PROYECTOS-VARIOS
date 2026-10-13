"""Pasa el sombreado del reporte Turnitin (PDF) al Word original.

Genera:
  - <salida>_sombreado.docx : copia del Word con cada coincidencia resaltada
    (un color de Word por fuente) para ubicarla y corregirla.
  - <salida>_coincidencias.xlsx : una fila por coincidencia con párrafo,
    sección, fuente, tipo de fuente, texto coincidente y % del párrafo afectado.

Uso:
    python herramientas/mapear_turnitin.py tesis.docx reporte_turnitin.pdf trabajos/x/salida/tesis
"""
import re
import sys
import unicodedata
from difflib import SequenceMatcher

import docx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from docx_lib import parrafos, resaltar_tramo
from turnitin_pdf import extraer

COLORES_WORD = ['yellow', 'green', 'magenta', 'red', 'blue', 'darkYellow', 'lightGray', 'darkCyan', 'darkMagenta']


def norm(s):
    s = unicodedata.normalize('NFKC', s).lower()
    s = re.sub(r'-\s+', '', s)  # guiones de corte de línea
    return re.sub(r'\s+', ' ', s).strip()


def _mapa_norm(t):
    """Texto normalizado + índice de cada carácter normalizado en el original."""
    out, idx, prev_esp = [], [], False
    for i, ch in enumerate(t):
        for c in unicodedata.normalize('NFKC', ch).lower():
            if c.isspace():
                if prev_esp or not out:
                    continue
                c, prev_esp = ' ', True
            else:
                prev_esp = False
            out.append(c)
            idx.append(i)
    return ''.join(out), idx


def ngramas(s, n=4):
    w = s.split()
    return {' '.join(w[i:i + n]) for i in range(max(0, len(w) - n + 1))}


def ubicar(frag, pars):
    """Devuelve (par, inicio, fin, similitud) del mejor tramo para el fragmento."""
    f = norm(frag)
    if len(f.split()) < 3:
        return None
    g = ngramas(f)
    mejor = None
    for p in pars:
        if not p['_ng'] or not (g & p['_ng']) and f not in p['_n']:
            continue
        k = p['_n'].find(f)
        if k >= 0:
            return p, p['_idx'][k], p['_idx'][k + len(f) - 1] + 1, 1.0
        sm = SequenceMatcher(None, p['_n'], f, autojunk=False)
        bloques = [b for b in sm.get_matching_blocks() if b.size >= 12]
        if not bloques:
            continue
        a = bloques[0].a
        b = bloques[-1].a + bloques[-1].size
        cubierto = sum(x.size for x in bloques) / len(f)
        if mejor is None or cubierto > mejor[3]:
            mejor = (p, p['_idx'][a], p['_idx'][b - 1] + 1, cubierto)
    if mejor and mejor[3] >= 0.5:
        return mejor
    return None


def main(ruta_docx, ruta_pdf, base_salida):
    rep = extraer(ruta_pdf)
    fuentes = {f['n']: f for f in rep['fuentes']}
    d = docx.Document(ruta_docx)
    pars = [p for p in parrafos(d) if p['texto'].strip()]
    for p in pars:
        p['_n'], p['_idx'] = _mapa_norm(p['texto'])
        p['_ng'] = ngramas(p['_n'])

    filas, no_ubicados, tramos = [], [], []
    for fr in rep['fragmentos']:
        r = ubicar(fr['texto'], pars)
        if not r:
            no_ubicados.append(fr)
            continue
        p, a, b, sim = r
        tramos.append((p, a, b, fr['fuente']))
        filas.append([p['i'], p['seccion'], 'tabla' if p['en_tabla'] else 'texto', fr['fuente'],
                      fuentes.get(fr['fuente'], {}).get('tipo', ''), fuentes.get(fr['fuente'], {}).get('nombre', ''),
                      fuentes.get(fr['fuente'], {}).get('porcentaje', ''), fr['pagina'], p['texto'][a:b],
                      round(100 * (b - a) / max(1, len(p['texto']))), p['texto']])

    # resaltar en orden inverso por párrafo para no desplazar offsets
    for p, a, b, fu in sorted(tramos, key=lambda x: (x[0]['i'], -x[1])):
        color = COLORES_WORD[(fu or 0) % len(COLORES_WORD)]
        resaltar_tramo(p['el'], a, b, color)
    d.save(base_salida + '_sombreado.docx')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Coincidencias'
    cab = ['Párrafo', 'Sección', 'Ubicación', 'Fuente N°', 'Tipo de fuente', 'Fuente', '% fuente', 'Pág. PDF',
           'Texto coincidente', '% del párrafo', 'Párrafo completo', 'Acción']
    ws.append(cab)
    for f in filas:
        ws.append([_seguro(x) for x in f] + [''])
    ws2 = wb.create_sheet('Fuentes')
    ws2.append(['N°', 'Tipo', 'Fuente', '%'])
    for f in rep['fuentes']:
        ws2.append([f['n'], f['tipo'], _seguro(f['nombre']), f['porcentaje']])
    ws3 = wb.create_sheet('No ubicados')
    ws3.append(['Pág. PDF', 'Fuente N°', 'Texto'])
    for fr in no_ubicados:
        ws3.append([fr['pagina'], fr['fuente'], _seguro(fr['texto'])])
    for hoja in wb:
        for c in hoja[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F4E78')
        hoja.freeze_panes = 'A2'
        for col in hoja.columns:
            for c in col[1:]:
                c.alignment = Alignment(wrap_text=True, vertical='top')
    for col, ancho in zip('ABCDEFGHIJKL', [8, 30, 9, 9, 16, 35, 9, 8, 60, 10, 80, 25]):
        ws.column_dimensions[col].width = ancho
    wb.save(base_salida + '_coincidencias.xlsx')
    print(f"fragmentos {len(rep['fragmentos'])} | ubicados {len(filas)} | no ubicados {len(no_ubicados)} | "
          f"párrafos afectados {len({f[0] for f in filas})}")
    print('->', base_salida + '_sombreado.docx')
    print('->', base_salida + '_coincidencias.xlsx')


def _seguro(v):
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@'):
        return "'" + v
    return v


if __name__ == '__main__':
    main(*sys.argv[1:4])
